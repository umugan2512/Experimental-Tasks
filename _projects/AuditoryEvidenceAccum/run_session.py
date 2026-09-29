# !/usr/bin/python3
# -*- coding: utf-8 -*-
"""
Standalone session runner/dashboard -- one widget to launch a wheel-shaping stage task, watch its
live progress (everything PyBpod itself currently shows for a running session -- the "Wheel
Position" plot, every `WheelShapingPlots` outcome panel, a camera snippet), and record
weight/baseline/notes directly into `training_log.xlsx`, without having the PyBpod GUI open at all.
Lives at this project's own top level (a sibling of `tasks/`/`records/`/`experiments/`), not inside
`records/`, since it's the thing you open to RUN a session, not a record-keeping script.

Run directly: `python run_session.py` (or double-click `Run Session.bat` -- see that file).
Deliberately standalone, same "bench tool, not a PyBpod GUI task" reasoning as
`Calibration/rig_check.py`/`calibrate_liquid.py`: it never opens a Bpod/rotary/HiFi connection of
its own, and launches a real stage task the exact same way `board_com.py:run_task()` does (create
the session folder, write a matching `user_settings.py`, `subprocess.Popen` the task script with
that folder as `cwd`) -- so from the task script's own point of view, nothing is different from
being launched by the real PyBpod GUI.

**Isolation**: this is the ONE new file. It only *reads* via `session_csv_parser.py`/
`build_training_log.py`/`wheel_shaping_plots.py` (imports them, never edits them) and only
*writes* to `training_log.xlsx`'s existing manual cells (the same cells a human already edits by
hand) and to a brand-new timestamped session folder it creates for a task it launches. No existing
task script or shared module is touched.

**Live view, everything reconstructed from the session's own CSV file (no live Bpod connection
needed), all embedded in ONE window (no separate popups)**:
- A small wheel-position-over-time plot, mirroring `pybpod-gui-plugin-rotaryencoder`'s own "Wheel
  Position" live plot -- built from `WHEEL_POS` VAL rows, which already encode both elapsed time
  and position together (`"<elapsed_s>,<position_deg>"`). Also overlays each trial's own threshold
  (a step line -- Stage 2's own staircase changes it trial to trial) and a light vertical marker at
  every trial's own start time.
- The FULL `WheelShapingPlots` panel set (raster, progress, accuracy, engagement, sidebias,
  percent_outcome, psychometric, reaction_time, reward_licks, lick_timeline) -- reused directly via
  its own `embed=True` mode (see `wheel_shaping_plots.py`), which builds a bare Figure/
  FigureCanvasQTAgg instead of letting `plt.subplot_mosaic()` open its own top-level window, so its
  canvas can be embedded directly into this window's own layout (inside a scroll area, since the
  full panel set is taller than most screens). Fed incrementally: each new trial found in the
  parsed CSV since the last poll has its `add_trial()` arguments reconstructed from the raw STATE/
  EVENT/VAL rows (`classify_trial()` for outcome/side, direct VAL lookups for threshold/gain/ITI/
  trial_type/response_time/click counts). `p_right_target` (Stage 3/4's debiasing target) is the
  one value never logged -- it's fully recoverable anyway, since `debiasing.next_side_after_error()`'s
  target is a deterministic function of only the PREVIOUS trial's own logged outcome/side, replayed
  here trial-by-trial rather than guessed.
- A camera snippet, refreshed from `CameraRecorder`'s own new live-frame JPEG export
  (`<video_stem>_latest.jpg`, see `_shared/camera_recorder.py`) rather than the growing
  `session_video.avi` file itself -- confirmed directly that `cv2.VideoWriter` (XVID/.avi) writes
  NOTHING to disk at all until `release()` is called (the file stays 0 bytes for the whole
  session), so reading the growing video file was never actually going to work; the small JPEG
  companion file is a plain, complete, always-readable image with no video-container concerns.

**Known, accepted approximations** (documented rather than silently wrong):
- `magnitude_deg` fed to the raster is reconstructed as "that trial's own threshold if the wheel
  crossed it, else 0.0" -- confirmed by reading all four real stage scripts directly, this is
  EXACTLY what they themselves already pass to `add_trial()` (none of them log the true continuous
  peak deflection), so this is not an approximation at all, just matching existing behavior.
- `lick_times_abs`/`reward_time_abs` are anchored to each trial's own logged absolute-time VAL
  (`TRIAL_START` for Stage 1/2, `SM_SEND_TIME` for Stage 3/4) plus that trial's raw event/state
  offsets. For Stage 3/4, a lick occurring during the earlier CUE period (before the choice-period
  state machine sends) is not separately recoverable this way and is simply not included in the
  lick timeline -- a minor, rare gap (most licks happen during/after the choice period).
"""
import contextlib
import datetime
import glob
import io
import json
import os
import re
import subprocess
import sys
import time
import traceback
import uuid

import cv2
import numpy as np

import matplotlib
matplotlib.use('Qt5Agg')   # same fix already applied to full_protocol_lookback_test.py -- one Qt
                            # event loop for the whole process, see CLAUDE.md's PyEval_RestoreThread
                            # crash note.

from PyQt5.QtCore import Qt, QTimer
from PyQt5.QtGui import QImage, QPixmap
from PyQt5.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QFormLayout, QComboBox, QPushButton, QLabel,
    QGroupBox, QMessageBox, QLineEdit, QTextEdit, QDoubleSpinBox, QCheckBox, QScrollArea,
    QDialog, QDialogButtonBox, QListWidget, QListWidgetItem)

_PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
_RECORDS_DIR = os.path.join(_PROJECT_DIR, 'records')
_TASKS_DIR = os.path.join(_PROJECT_DIR, 'tasks')
_SHARED_DIR = os.path.join(_TASKS_DIR, '_wheel_shaping_shared')
_OUTPUT_XLSX = os.path.join(_RECORDS_DIR, 'training_log.xlsx')

sys.path.insert(0, _SHARED_DIR)
sys.path.insert(0, _RECORDS_DIR)   # build_training_log.py lives in records/, a sibling directory
                                    # now that this script lives at the project's own top level
                                    # rather than inside records/ itself -- an explicit path entry
                                    # instead of relying on the script's own directory implicitly
                                    # being on sys.path (which only happens for the __main__ entry
                                    # script, not a module imported from elsewhere).
import session_csv_parser     # noqa: E402
import build_training_log     # noqa: E402
import wheel_shaping_plots    # noqa: E402
import debiasing               # noqa: E402

import openpyxl  # noqa: E402

VAR_BPOD_FIRMWARE_VERSION = '22'   # matches every real session-local user_settings.py's own
                                    # TARGET_BPOD_FIRMWARE_VERSION on this rig -- same "machine-
                                    # specific constant" convention as VAR_BPOD_SERIAL_PORT
                                    # elsewhere in this codebase, see pybpodapi/settings.py's own
                                    # module-level default for confirmation.
VAR_PROJECT_NAME = 'AuditoryEvidenceAccum'

# The interpreter used to launch a TASK subprocess, deliberately separate from whatever interpreter
# is running this dashboard's own GUI (sys.executable) -- run_session.py's own deps (PyQt5/cv2/
# matplotlib/openpyxl) and a task script's own deps (pybpodapi + hardware modules) are genuinely
# different, so forcing them to share one interpreter only worked by coincidence when both happened
# to be the same env. Confirmed as a real, already-hit failure: a task launched with the wrong
# interpreter dies instantly with "ModuleNotFoundError: No module named 'pybpodapi'", and previously
# that was only ever visible in a log file inside the new session folder, not in the UI at all (see
# SessionRunnerWindow's startup self-check / post-launch quick-death detection below).
# Machine-specific -- edit/extend per box, same convention as VAR_BPOD_SERIAL_PORT elsewhere in
# this codebase. Tries each known real environment in order (the real rig's own conda env is named
# "pybpod-environment"; this dev machine's is "UM-pybpod", confirmed directly:
# `C:\Users\Gil\anaconda3\envs\UM-pybpod\python.exe -c "import pybpodapi"` succeeds, resolving to
# this repo's own editable pybpod-api install), falling back to sys.executable only if none exist.
_TASK_PYTHON_CANDIDATES = [
    r'C:\Users\2P-Behav\.conda\envs\pybpod-environment\python.exe',
    r'C:\Users\Gil\anaconda3\envs\UM-pybpod\python.exe',
]
VAR_TASK_PYTHON_EXE = next((p for p in _TASK_PYTHON_CANDIDATES if os.path.exists(p)), sys.executable)
VAR_THRESHOLD_FINAL_DEG = 35.0     # matches every stage script's own VAR_THRESHOLD_FINAL_DEG /
                                    # build_training_log.py's own hardcoded threshold_final_deg.

VAR_SESSION_POLL_MS = 1000    # how often to look for a newer session folder on disk
VAR_PLOTS_POLL_MS = 2000      # how often to re-parse the watched CSV and feed new trials
VAR_CAMERA_POLL_MS = 1000     # how often to read the latest camera JPEG (see grab_latest_frame())
VAR_STOP_GRACE_S = 8.0        # how long to wait for a graceful 'close' before escalating to kill
VAR_QUICK_DEATH_CHECK_MS = 2000   # how long after launch to check whether the task subprocess has
                                   # already died (e.g. missing dependency, bad COM port) -- long
                                   # enough for a real Bpod handshake/import to get underway, short
                                   # enough that a genuine crash is still reported promptly.

_VALID_NAME_RE = re.compile(r'^[A-Za-z0-9_-]+$')


def task_python_can_import_pybpodapi():
    """ Startup self-check: does VAR_TASK_PYTHON_EXE actually have pybpodapi importable? Run once
    at window construction so a misconfigured/missing task environment is reported immediately as
    a visible warning banner, instead of only surfacing after a session folder has already been
    created and the task subprocess has already died silently (the exact failure mode that
    prompted this check -- see VAR_TASK_PYTHON_EXE's own docstring above). """
    try:
        # stdout=/stderr=PIPE, not capture_output=True -- this whole file needs to run under
        # Python 3.6 (e.g. launched via the "UM-pybpod" Anaconda env, which already bundles every
        # GUI dependency this file needs alongside pybpodapi itself), and capture_output is a
        # Python 3.7+-only subprocess.run() kwarg. Confirmed as a real, already-hit bug: under
        # Python 3.6 this raised TypeError on every call, silently swallowed by the broad except
        # below, making the self-check permanently (and wrongly) report "not importable".
        result = subprocess.run([VAR_TASK_PYTHON_EXE, '-c', 'import pybpodapi'],
                                 stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=15)
        return result.returncode == 0
    except Exception:
        return False


_DASHBOARD_PROTOCOL_INFO = {
    'stage1_wheel_shaping': {
        'stage': 1,
        'anchor_key': 'TRIAL_START',
        'threshold_key': 'THRESHOLD_DEG',
        'outcome_map': {'rewarded': 'Rewarded', 'no_movement': 'NoMovement'},
        'crossed_outcomes': {'rewarded'},
    },
    'stage2_threshold_staircase': {
        'stage': 2,
        'anchor_key': 'TRIAL_START',
        'threshold_key': 'THRESHOLD_DEG',
        'outcome_map': {'rewarded': 'Rewarded', 'withheld': 'Withheld', 'no_movement': 'NoMovement'},
        'crossed_outcomes': {'rewarded', 'withheld'},
    },
    'stage3_clicks_direction': {
        'stage': 3,
        'anchor_key': 'SM_SEND_TIME',
        'threshold_key': 'RESPONSE_THRESHOLD_DEG',
        'outcome_map': {'rewarded': 'Reward', 'incorrect': 'NoReward', 'no_movement': 'NoResponse',
                         'aborted': 'Abort'},
        'crossed_outcomes': {'rewarded', 'incorrect'},
    },
    'stage4_resume_staircases': {
        'stage': 4,
        'anchor_key': 'SM_SEND_TIME',
        'threshold_key': 'RESPONSE_THRESHOLD_DEG',
        'outcome_map': {'rewarded': 'Reward', 'incorrect': 'NoReward', 'no_movement': 'NoResponse',
                         'aborted': 'Abort'},
        'crossed_outcomes': {'rewarded', 'incorrect'},
    },
}

_OUTCOME_STATE_NAME = {
    'stage1_wheel_shaping': {'rewarded': 'Reward'},
    'stage2_threshold_staircase': {'rewarded': ('RewardL', 'RewardR'),
                                    'withheld': ('NoRewardL', 'NoRewardR')},
    'stage3_clicks_direction': {'rewarded': 'Reward', 'incorrect': 'ErrorConsumption'},
    'stage4_resume_staircases': {'rewarded': 'Reward', 'incorrect': 'ErrorConsumption'},
}

_USER_SETTINGS_TEMPLATE = """SETTINGS_PRIORITY = 0

PYBPOD_SERIAL_PORT       = '{serial_port}'
PYBPOD_NET_PORT          = {net_port}

{bnc_ports_line}
{wired_ports_line}
{behavior_ports_line}

PYBPOD_API_LOG_LEVEL = None
PYBPOD_API_LOG_FILE  = None

PYBPOD_API_STREAM2STDOUT = True
PYBPOD_API_ACCEPT_STDIN  = True

PYBPOD_PROTOCOL \t= '{protocol}'
PYBPOD_CREATOR \t\t= '{creator}'
PYBPOD_PROJECT \t\t= '{project}'
PYBPOD_EXPERIMENT \t= '{experiment}'
PYBPOD_BOARD \t\t= '{board}'
PYBPOD_SETUP \t\t= '{setup}'
PYBPOD_SESSION \t\t= '{session}'
PYBPOD_SESSION_PATH = '{session_path}'
PYBPOD_SUBJECTS \t= [{subjects}]
PYBPOD_SUBJECT_EXTRA = ''
PYBPOD_USER_EXTRA = ''


TARGET_BPOD_FIRMWARE_VERSION = '{firmware}'

#import logging
#PYBPOD_API_LOG_LEVEL = logging.DEBUG
#PYBPOD_API_LOG_FILE  = 'pybpod-api.log'

PYBPOD_VARSNAMES = []
"""


# --- project/board/subject/user discovery (mirrors what the PyBpod GUI itself already has on
# disk -- never guessed/hardcoded) ------------------------------------------------------------

def discover_stage_configs():
    """ Returns a list of {experiment, setup, board, task_name, task_script, subjects}, one per
    already-GUI-configured experiment/setup this project has on disk (scanned, not a hardcoded
    table, so a not-yet-created stage -- e.g. Stage 4 before its first GUI setup -- simply doesn't
    appear rather than needing separate scaffolding-creation logic this tool doesn't own). Only
    setups whose task actually has a real script under tasks/ are included. """
    results = []
    pattern = os.path.join(_PROJECT_DIR, 'experiments', '*', 'setups', '*', '*.json')
    for path in sorted(glob.glob(pattern)):
        try:
            with open(path, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception:
            continue
        setup_dir = os.path.dirname(path)
        experiment_dir = os.path.dirname(os.path.dirname(setup_dir))
        experiment = os.path.basename(experiment_dir)
        setup = os.path.basename(setup_dir)
        task_name = data.get('task')
        task_script = os.path.join(_TASKS_DIR, task_name or '', (task_name or '') + '.py')
        if not task_name or not os.path.exists(task_script):
            continue
        results.append({
            'experiment': experiment, 'setup': setup, 'board': data.get('board'),
            'task_name': task_name, 'task_script': task_script,
            'subjects': list(data.get('subjects') or []),
        })
    results.sort(key=lambda c: build_training_log._stage_sort_key(c['task_name']))
    return results


def discover_users():
    results = []
    for path in sorted(glob.glob(os.path.join(_PROJECT_DIR, 'users', '*', '*.json'))):
        try:
            with open(path, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception:
            continue
        name = os.path.splitext(os.path.basename(path))[0]
        uuid_ = data.get('__UUID4__')
        if uuid_:
            results.append({'name': name, 'uuid': uuid_})
    return results


def load_subject_uuid(subject_name):
    path = os.path.join(_PROJECT_DIR, 'subjects', subject_name, subject_name + '.json')
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)['__UUID4__']


def load_board_config(board_name):
    path = os.path.join(_PROJECT_DIR, 'boards', board_name, board_name + '.json')
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def discover_subjects():
    """ Returns [{'name', 'uuid'}] for every subjects/*/*.json on disk, regardless of whether it's
    attached to a Setup yet -- used by the New Setup dialog's subject checklist (discover_stage_
    configs()'s own per-stage subject lists only include already-attached subjects). """
    results = []
    for path in sorted(glob.glob(os.path.join(_PROJECT_DIR, 'subjects', '*', '*.json'))):
        try:
            with open(path, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception:
            continue
        name = os.path.splitext(os.path.basename(path))[0]
        results.append({'name': name, 'uuid': data.get('__UUID4__')})
    return results


def discover_boards():
    results = []
    for path in sorted(glob.glob(os.path.join(_PROJECT_DIR, 'boards', '*', '*.json'))):
        try:
            with open(path, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception:
            continue
        name = os.path.splitext(os.path.basename(path))[0]
        results.append({'name': name, 'uuid': data.get('__UUID4__')})
    return results


def discover_all_tasks():
    """ Every task script this project has under tasks/<name>/<name>.py -- for the New Setup
    dialog's Task dropdown (discover_stage_configs() only lists tasks an existing Setup already
    uses, not the raw set of available scripts, since a new Setup needs to be able to pick one that
    isn't in use anywhere yet, e.g. stage4_resume_staircases before its first Setup exists). """
    results = []
    for path in sorted(glob.glob(os.path.join(_TASKS_DIR, '*', '*.py'))):
        name = os.path.splitext(os.path.basename(path))[0]
        if name == os.path.basename(os.path.dirname(path)) and not name.startswith('_'):
            results.append(name)
    return results


def _validate_new_name(name):
    if not name or not _VALID_NAME_RE.match(name):
        raise ValueError("Name must be non-empty and contain only letters, numbers, '_' or '-'.")


def _scadict_boilerplate(def_text):
    """ Matches every real project JSON's own boilerplate shape exactly -- confirmed against
    `sca/formats/json.py`'s `scadict`/`dump()` (isoformat(' ') timestamps, this fixed key set) and
    `sca/format_flags.py`'s own flag names, by reading PyBpod's actual source. """
    now = datetime.datetime.now().isoformat(' ')
    return {
        '__UUID4__': str(uuid.uuid4()),
        '__CREATED-ON__': now,
        '__UPDATED-ON__': now,
        '__SOFTWARE__': 'PyBpod GUI API v1.8.2',
        '__DEF-URL__': 'http://pybpod.readthedocs.org',
        '__DEF-TEXT__': def_text,
    }


def _write_pybpod_json(path, data):
    # indent=4, sort_keys=True -- matches sca/formats/json.py's own dump(), confirmed against
    # every real example file already in this project (alphabetically-sorted keys).
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=4, sort_keys=True)


def create_subject(name):
    """ Writes subjects/<name>/<name>.json, matching pybpod-gui-api's own SubjectIO.save() shape
    exactly -- unassigned to any Setup yet. Note "setup": "None" is literally the STRING "None",
    not JSON null -- confirmed from the real source (`str(self.setup.uuid4 if self.setup else
    None)`), reproduced exactly rather than "fixed" to null, so a real PyBpod GUI opening this
    project later sees byte-for-byte what it would have written itself. """
    _validate_new_name(name)
    subject_dir = os.path.join(_PROJECT_DIR, 'subjects', name)
    if os.path.exists(subject_dir):
        raise ValueError('Subject "{0}" already exists.'.format(name))
    os.makedirs(subject_dir)
    data = _scadict_boilerplate('This file contains information about a subject used on PyBpod GUI.')
    data['setup'] = 'None'
    _write_pybpod_json(os.path.join(subject_dir, name + '.json'), data)
    return data['__UUID4__']


def create_setup(name, board_name, task_name, subject_names):
    """ Creates experiments/<name>/setups/<name>/ and writes <name>.json matching pybpod-gui-api's
    own SetupBaseIO.save() shape exactly: board/task as plain name strings (not UUIDs), subjects as
    a name list, __EXTERNAL-REF__ containing the board's UUID plus every assigned subject's UUID
    (add_external_ref()). `name` is used for BOTH the experiment and setup directory -- matches
    this project's own established 1:1 convention (every existing Stage1/2/3 example does this;
    confirmed no experiment-level JSON is ever written at all -- an "experiment" is purely a
    directory). Also updates each assigned subject's own JSON "setup" field to this new setup's
    UUID -- the two-way link the real PyBpod GUI itself maintains. """
    _validate_new_name(name)
    setup_dir = os.path.join(_PROJECT_DIR, 'experiments', name, 'setups', name)
    if os.path.exists(setup_dir):
        raise ValueError('An experiment/setup named "{0}" already exists.'.format(name))

    board_cfg_path = os.path.join(_PROJECT_DIR, 'boards', board_name, board_name + '.json')
    with open(board_cfg_path, 'r', encoding='utf-8') as f:
        board_uuid = json.load(f)['__UUID4__']

    subject_paths = {n: os.path.join(_PROJECT_DIR, 'subjects', n, n + '.json') for n in subject_names}
    subject_uuids = {}
    for subject_name, subject_path in subject_paths.items():
        with open(subject_path, 'r', encoding='utf-8') as f:
            subject_uuids[subject_name] = json.load(f)['__UUID4__']

    os.makedirs(setup_dir)
    data = _scadict_boilerplate(
        'This file contains information about a PyBpod experiment setup.')
    data['board'] = board_name
    data['task'] = task_name
    data['subjects'] = list(subject_names)
    data['detached'] = False
    data['variables'] = []
    data['update-variables'] = False
    data['__EXTERNAL-REF__'] = [board_uuid] + [subject_uuids[n] for n in subject_names]
    _write_pybpod_json(os.path.join(setup_dir, name + '.json'), data)

    setup_uuid = data['__UUID4__']
    for subject_name, subject_path in subject_paths.items():
        with open(subject_path, 'r', encoding='utf-8') as f:
            subject_data = json.load(f)
        subject_data['setup'] = setup_uuid
        subject_data['__UPDATED-ON__'] = datetime.datetime.now().isoformat(' ')
        _write_pybpod_json(subject_path, subject_data)

    return setup_uuid


# --- session-folder creation + subprocess launch, mirroring board_com.py:run_task() exactly ----

def build_user_settings_content(board_name, board_cfg, experiment, setup, session_name,
                                 session_path, task_name, subject_name, subject_uuid, user_name,
                                 user_uuid):
    bnc = board_cfg.get('enabled-bncports')
    wired = board_cfg.get('enabled-wiredports')
    behavior = board_cfg.get('enabled-behaviorports')
    # PYBPOD_SUBJECTS uses plain Python str() (single-quoted); PYBPOD_CREATOR uses json.dumps()
    # (double-quoted) -- confirmed against board_com.py's own run_task(), which builds them via
    # two genuinely different serializations (str([s.name, str(s.uuid4)]) vs.
    # json.dumps([user.name, str(user.uuid4), user.connection])), not a copy-paste inconsistency.
    subjects_field = '"{0}"'.format(str([subject_name, subject_uuid]))
    creator_field = json.dumps([user_name, user_uuid, 'local'])
    session_path_field = os.path.abspath(session_path).encode('unicode_escape').decode()
    return _USER_SETTINGS_TEMPLATE.format(
        serial_port=board_cfg.get('serial-port'), net_port=board_cfg.get('net-port'),
        bnc_ports_line=('BPOD_BNC_PORTS_ENABLED = {0}'.format(bnc) if bnc else ''),
        wired_ports_line=('BPOD_WIRED_PORTS_ENABLED = {0}'.format(wired) if wired else ''),
        behavior_ports_line=('BPOD_BEHAVIOR_PORTS_ENABLED = {0}'.format(behavior)
                              if behavior else ''),
        protocol=task_name, creator=creator_field, project=VAR_PROJECT_NAME, experiment=experiment,
        board=board_name, setup=setup, session=session_name, session_path=session_path_field,
        subjects=subjects_field, firmware=VAR_BPOD_FIRMWARE_VERSION)


class TaskLauncher(object):
    """ Creates a brand-new session folder + user_settings.py and launches the task subprocess,
    the exact same way board_com.py:run_task() does -- so the launched task script sees nothing
    different from being started by the real PyBpod GUI. Owns the Popen handle for Stop/Kill. """

    def __init__(self):
        self.proc = None
        self.session_path = None

    @property
    def is_running(self):
        return self.proc is not None and self.proc.poll() is None

    def start(self, stage_config, subject_name, user):
        if self.is_running:
            raise RuntimeError('A task launched from this dashboard is already running.')

        board_cfg = load_board_config(stage_config['board'])
        subject_uuid = load_subject_uuid(subject_name)

        session_name = datetime.datetime.now().strftime('%Y%m%d-%H%M%S')
        session_path = os.path.join(_PROJECT_DIR, 'experiments', stage_config['experiment'],
                                     'setups', stage_config['setup'], 'sessions', session_name)
        os.makedirs(session_path, exist_ok=True)

        with open(os.path.join(session_path, '__init__.py'), 'w') as f:
            pass

        settings_content = build_user_settings_content(
            stage_config['board'], board_cfg, stage_config['experiment'], stage_config['setup'],
            session_name, session_path, stage_config['task_name'], subject_name, subject_uuid,
            user['name'], user['uuid'])
        with open(os.path.join(session_path, 'user_settings.py'), 'w') as f:
            f.write(settings_content)

        env = os.environ.copy()
        # Only the session folder itself -- NOT this (parent) process's own sys.path.
        # board_com.py's real run_task() does append its own sys.path here too, but that's only
        # safe there because the launcher and the launched task always share the SAME interpreter.
        # This dashboard deliberately launches the task under a DIFFERENT interpreter
        # (VAR_TASK_PYTHON_EXE) than whatever runs the GUI itself -- forwarding THIS process's own
        # sys.path (e.g. a Python 3.9 install's own stdlib paths) onto a Python 3.6 child's
        # PYTHONPATH corrupts its stdlib resolution entirely. Confirmed directly: doing so produced
        # "Fatal Python error: Py_Initialize: can't initialize sys standard streams" /
        # "No module named '_abc'" from the child trying to load Python 3.9's io.py/abc.py under
        # Python 3.6. The session folder alone is all `import user_settings` inside pybpodapi
        # itself actually needs.
        env['PYTHONPATH'] = os.path.abspath(session_path)
        # Tells the task script (stage1-4, see their own identical comment at their own
        # WheelShapingPlots(...) construction) to suppress its own separate live-plot popup window
        # -- this dashboard's own embedded view already shows the same data, reconstructed from the
        # session CSV. Absent entirely for a real PyBpod-GUI-launched session (board_com.py never
        # sets this), so that path is completely unaffected.
        env['RUN_SESSION_DASHBOARD'] = '1'

        stdout_log = open(os.path.join(session_path, 'dashboard_launch_stdout.log'), 'w')
        stderr_log = open(os.path.join(session_path, 'dashboard_launch_stderr.log'), 'w')
        self.proc = subprocess.Popen(
            [VAR_TASK_PYTHON_EXE, os.path.abspath(stage_config['task_script'])],
            stdin=subprocess.PIPE, stdout=stdout_log, stderr=stderr_log, cwd=session_path, env=env)
        self.session_path = session_path
        return session_path

    def stop(self, graceful=True):
        """ Graceful: writes 'close\\r\\n' to stdin -- the exact same command board_com.py's own
        stop_task() sends, causing the task's current Bpod call to unwind and its normal
        post-loop cleanup (struct export, Bpod.close() writing SESSION-ENDED) to still run. Caller
        is responsible for escalating to terminate()/kill() if the process hasn't exited after
        VAR_STOP_GRACE_S (see SessionRunnerWindow._on_stop_check). """
        if self.proc is None or self.proc.poll() is not None:
            return
        try:
            if graceful and self.proc.stdin is not None:
                self.proc.stdin.write('close\r\n'.encode())
                self.proc.stdin.flush()
            else:
                self.proc.terminate()
        except Exception:
            print(traceback.format_exc(), flush=True)

    def force_kill(self):
        if self.proc is not None and self.proc.poll() is None:
            try:
                self.proc.kill()
            except Exception:
                print(traceback.format_exc(), flush=True)


# --- CSV -> WheelShapingPlots.add_trial() reconstruction ---------------------------------------

def _val_float(trial, key, session_vals=None):
    """ Checks this trial's own vals first, falling back to session_vals -- some registrations
    (e.g. Stage 1's own THRESHOLD_DEG, which only changes across SESSIONS not within one) happen
    once before the first TRIAL marker rather than every trial, landing in session_vals instead of
    any individual trial's own vals (see session_csv_parser.find_val_backward()'s own docstring for
    the same distinction). """
    v = trial['vals'].get(key)
    if v is None and session_vals is not None:
        v = session_vals.get(key)
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _val_str(trial, key, session_vals=None):
    v = trial['vals'].get(key)
    if v is None and session_vals is not None:
        v = session_vals.get(key)
    return v


class PlotFeeder(object):
    """ Owns one long-lived WheelShapingPlots instance for the currently-watched session and feeds
    it incrementally -- each poll() call only replays trials beyond the count already fed, mirroring
    how a real stage script itself calls add_trial() once per trial rather than replaying the whole
    session from scratch every tick. """

    def __init__(self, protocol_name, first_trial_vals, session_vals):
        info = _DASHBOARD_PROTOCOL_INFO[protocol_name]
        self.protocol_name = protocol_name
        self.info = info
        self.session_vals = session_vals
        prev_session_values = {}
        starting_threshold = (first_trial_vals.get(info['threshold_key'])
                               or session_vals.get(info['threshold_key']))
        try:
            starting_threshold = float(starting_threshold) if starting_threshold is not None else None
        except (TypeError, ValueError):
            starting_threshold = None
        if starting_threshold is not None:
            prev_session_values['threshold_deg'] = starting_threshold
        session_status = {}
        cue_abort = (first_trial_vals.get('CUE_ABORT_THRESHOLD_DEG')
                     or session_vals.get('CUE_ABORT_THRESHOLD_DEG'))
        if cue_abort is not None:
            try:
                session_status['in_trial_threshold_deg'] = float(cue_abort)
            except (TypeError, ValueError):
                pass
        self.plots = wheel_shaping_plots.WheelShapingPlots(
            info['stage'], VAR_THRESHOLD_FINAL_DEG, prev_session_values=prev_session_values,
            session_status=session_status, embed=True)
        self._n_fed = 0
        self._replay_prev_outcome = None
        self._replay_prev_side = None

    def feed_new_trials(self, trials, session_ended=False):
        """ Holds back the LAST entry in `trials` -- this dashboard's own file read races against
        the task subprocess's own burst of STATE/EVENT/VAL row writes for whichever trial is
        currently most recent (each trial's own rows are flushed as a tight burst, not atomically
        -- confirmed directly by reading pybpod-api's own session.py: `Session.__add__()` flushes
        after every single row, individually). A poll landing mid-burst can see that trial with a
        partial states/vals dict and misclassify it (typically 'unknown', silently skipped by
        _reconstruct_kwargs()) -- and since `self._n_fed` is a COUNT, not a per-trial identity
        check, that trial would otherwise be marked "already fed" forever and never reconsidered
        even once its real data lands on a later poll. Holding back the tail entry costs one poll
        cycle (~2s) of display lag and fixes this permanently -- confirmed by a dedicated test
        feeding the SAME trial index twice, first incomplete then complete.

        `session_ended=True` (the session's own SESSION-ENDED INFO row is present, i.e. Bpod.close()
        already ran and the file is finalized -- no more trials are ever coming) skips holding
        anything back, since there's no "next trial" left to ever supersede a final held-back one --
        without this, the very last trial of a finished session would stay permanently unfed. """
        stable_trials = trials if session_ended or not trials else trials[:-1]
        for trial in stable_trials[self._n_fed:]:
            kwargs = self._reconstruct_kwargs(trial)
            if kwargs is not None:
                self.plots.add_trial(**kwargs)
        self._n_fed = len(stable_trials)

    def _reconstruct_kwargs(self, trial):
        info = self.info
        config = session_csv_parser.PROTOCOL_CONFIG[self.protocol_name]
        outcome_raw, event_side, _reward_duration, _consumed = session_csv_parser.classify_trial(
            trial, config)
        side = _val_str(trial, 'TRIAL_SIDE', self.session_vals) or event_side or 'R'
        outcome = info['outcome_map'].get(outcome_raw)
        if outcome is None:
            return None   # 'unknown' -- no recognizable outcome states visited, nothing to plot

        threshold_deg = _val_float(trial, info['threshold_key'], self.session_vals)
        magnitude_deg = (threshold_deg or 0.0) if outcome_raw in info['crossed_outcomes'] else 0.0
        anchor = _val_float(trial, info['anchor_key'], self.session_vals)

        kwargs = dict(side=side, magnitude_deg=magnitude_deg, threshold_deg=threshold_deg or 0.0,
                      outcome=outcome)

        lick_times_abs = []
        if anchor is not None:
            lick_times_abs = [anchor + t for t in trial['events'].get('Port1In', [])]
        kwargs['lick_times_abs'] = lick_times_abs

        reward_state_name = _OUTCOME_STATE_NAME.get(self.protocol_name, {}).get('rewarded')
        if outcome_raw == 'rewarded' and anchor is not None:
            name = reward_state_name
            if isinstance(name, tuple):
                name = 'RewardL' if side == 'L' else 'RewardR'
            state = trial['states'].get(name)
            if state is not None and state[0] is not None:
                kwargs['reward_time_abs'] = anchor + state[0]

        if info['stage'] == 1:
            kwargs['gain_mult'] = _val_float(trial, 'GAIN_MULT', self.session_vals)
        elif info['stage'] == 2:
            kwargs['direction_ratio'] = _val_float(trial, 'DIRECTION_RATIO', self.session_vals)
            iti_state = trial['states'].get('ITI')
            if iti_state is not None and iti_state[2] is not None:
                kwargs['iti_s'] = iti_state[2]
        else:   # stage 3/4
            trial_type = _val_str(trial, 'TRIAL_TYPE', self.session_vals)
            kwargs['trial_type'] = trial_type
            in_trial_threshold = _val_float(trial, 'CUE_ABORT_THRESHOLD_DEG', self.session_vals)
            if in_trial_threshold is not None:
                kwargs['in_trial_threshold_deg'] = in_trial_threshold

            if trial_type == 'repeat':
                p_right_target = 1.0 if side == 'R' else 0.0
            elif self._replay_prev_outcome == 'incorrect':
                p_right_target = (debiasing.VAR_DEBIAS_REPEAT_PROB if self._replay_prev_side == 'R'
                                   else 1.0 - debiasing.VAR_DEBIAS_REPEAT_PROB)
            else:
                p_right_target = 0.5
            kwargs['p_right_target'] = p_right_target

            if outcome_raw in ('rewarded', 'incorrect'):
                outcome_state_name = _OUTCOME_STATE_NAME[self.protocol_name][outcome_raw]
                state = trial['states'].get(outcome_state_name)
                if state is not None and state[0] is not None:
                    kwargs['response_time_s'] = state[0]
                n_l = _val_float(trial, 'N_CLICKS_L', self.session_vals)
                n_r = _val_float(trial, 'N_CLICKS_R', self.session_vals)
                if n_l is not None and n_r is not None:
                    kwargs['click_diff'] = n_r - n_l
                self._replay_prev_outcome = 'correct' if outcome_raw == 'rewarded' else 'incorrect'
                self._replay_prev_side = side
            elif outcome_raw == 'no_movement':
                self._replay_prev_outcome = None
                self._replay_prev_side = side
            # outcome_raw == 'aborted': prev_outcome/prev_side deliberately left unchanged --
            # matches the real task scripts' own `continue` skipping that update for an Abort.

        return kwargs


def parse_wheel_position_series(trials, decimate_s=0.5):
    """ Every WHEEL_POS VAL across all trials, as (elapsed_s, position_deg) pairs -- WHEEL_POS is
    never collapsed to the last sample per trial (see session_csv_parser.STREAM_KEYS), so the full
    stream survives. Decimated to at most one sample per `decimate_s` seconds, same principle as
    the real "Wheel Position" plot's own DISPLAY_SAMPLE_INTERVAL_S, so a long session doesn't force
    an ever-growing full-resolution redraw every poll. """
    xs, ys = [], []
    last_t = None
    for trial in trials:
        raw = trial['vals'].get('WHEEL_POS')
        if not raw:
            continue
        for entry in raw:
            try:
                t_str, pos_str = entry.split(',')
                t, pos = float(t_str), float(pos_str)
            except (ValueError, AttributeError):
                continue
            if last_t is not None and (t - last_t) < decimate_s:
                continue
            xs.append(t)
            ys.append(pos)
            last_t = t
    return xs, ys


def build_trial_markers(trials, protocol_name, session_vals):
    """ For the wheel-position plot's overlay: returns (trial_start_times, threshold_segments).
    trial_start_times is a sorted list of each trial's own anchor time (a light vertical line per
    trial). threshold_segments is a list of (x_start, x_end, threshold_deg) -- one horizontal
    segment per trial, spanning from that trial's own start to the next trial's start (or a short
    pad past the last one), so a per-trial threshold CHANGE (e.g. Stage 2's own staircase) shows as
    a literal step rather than one flat line silently averaging over the whole session. Threshold
    falls back to the session-level VAL the same way _val_float() already does (e.g. Stage 1's own
    THRESHOLD_DEG, constant for the whole session, registered once before any trial). """
    if protocol_name not in _DASHBOARD_PROTOCOL_INFO:
        return [], []
    info = _DASHBOARD_PROTOCOL_INFO[protocol_name]
    points = []
    for trial in trials:
        anchor = _val_float(trial, info['anchor_key'], session_vals)
        if anchor is None:
            continue
        threshold = _val_float(trial, info['threshold_key'], session_vals)
        points.append((anchor, threshold))
    points.sort(key=lambda p: p[0])
    starts = [x for x, _threshold in points]
    segments = []
    for i, (x, threshold) in enumerate(points):
        if threshold is None:
            continue
        x_end = points[i + 1][0] if i + 1 < len(points) else x + 5.0
        segments.append((x, x_end, threshold))
    return starts, segments


# --- weight / notes / baseline write path, reusing build_training_log.py's own column layout ---

def find_row_for_session(ws, table_start_row, col_map, session_started):
    key_col = col_map.get('session_started')
    if key_col is None:
        return None
    for row in range(table_start_row + 1, ws.max_row + 1):
        raw = ws.cell(row=row, column=key_col).value
        if not raw:
            continue
        members = {k.strip() for k in str(raw).split(';') if k.strip()}
        if session_started in members:
            return row
    return None


def save_manual_fields(subject, session_started, weight_g, weight_after_g, hand_watered, notes,
                        header_fields):
    """ Regenerates training_log.xlsx first (so this session's own row definitely exists / is
    current -- same call "Update Training Log" already makes), then writes directly into the
    manual cells for that row plus the subject's header block (DOB/Sex/Strain/Baseline weight).
    Returns True if the session's own row was found and written. """
    build_training_log.build_workbook()
    wb = openpyxl.load_workbook(_OUTPUT_XLSX)
    ws, table_start_row = build_training_log._get_or_create_subject_sheet(wb, subject)
    col_map = build_training_log._scan_existing_column_labels(ws, table_start_row)

    for i, value in enumerate(header_fields):
        if value:
            ws.cell(row=2 + i, column=2, value=value)

    row = find_row_for_session(ws, table_start_row, col_map, session_started)
    found = row is not None
    if found:
        if weight_g is not None:
            ws.cell(row=row, column=col_map['weight_g'], value=weight_g)
        if weight_after_g is not None:
            ws.cell(row=row, column=col_map['weight_after_task_g'], value=weight_after_g)
        ws.cell(row=row, column=col_map['hand_watered'], value=hand_watered)
        if notes:
            ws.cell(row=row, column=col_map['notes'], value=notes)

    wb.save(_OUTPUT_XLSX)
    return found


def load_header_fields(subject):
    if not os.path.exists(_OUTPUT_XLSX):
        return ['', '', '', '']
    wb = openpyxl.load_workbook(_OUTPUT_XLSX)
    sheet_name = build_training_log._sheet_name_for(subject)
    if sheet_name not in wb.sheetnames:
        return ['', '', '', '']
    ws = wb[sheet_name]
    return [ws.cell(row=2 + i, column=2).value or '' for i in range(4)]


# --- camera snippet: reads CameraRecorder's own live-frame JPEG export, never the growing video
# file itself (confirmed cv2.VideoWriter writes nothing to disk at all until release() -- see
# camera_recorder.py's own "Live-frame JPEG export" docstring section) and never opens the camera
# device itself (that's exclusively owned by the running task's own CameraRecorder) ---------------

def grab_latest_frame(session_dir):
    """ Reads <video_stem>_latest.jpg -- a plain, complete, always-fully-written image file
    (overwritten in place by the running task's own CameraRecorder, ~5Hz) -- so this is a normal,
    reliable file read, not a video-container race. Returns None if the session hasn't started
    writing camera frames yet (no camera on this task, or too early in the session). """
    jpg_path = os.path.join(session_dir, 'session_video_latest.jpg')
    if not os.path.exists(jpg_path):
        return None
    return cv2.imread(jpg_path)


# --- newest-session discovery (same glob scan_all_sessions() already uses) ---------------------

def find_newest_session_csv():
    pattern = os.path.join(_PROJECT_DIR, 'experiments', '*', 'setups', '*', 'sessions', '*', '*.csv')
    paths = glob.glob(pattern)
    if not paths:
        return None
    return max(paths, key=os.path.getmtime)


# --- New Subject / New Setup dialogs, mirroring what the PyBpod GUI's own equivalent actions
# would ask for -------------------------------------------------------------------------------

class NewSubjectDialog(QDialog):
    def __init__(self, parent=None):
        super(NewSubjectDialog, self).__init__(parent)
        self.setWindowTitle('New Subject')
        layout = QFormLayout()
        self.name_input = QLineEdit()
        layout.addRow('Subject name:', self.name_input)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addRow(buttons)
        self.setLayout(layout)

    def subject_name(self):
        return self.name_input.text().strip()


class NewSetupDialog(QDialog):
    def __init__(self, boards, tasks, subjects, parent=None):
        super(NewSetupDialog, self).__init__(parent)
        self.setWindowTitle('New Setup')
        layout = QFormLayout()
        self.name_input = QLineEdit()
        layout.addRow('Name (experiment + setup):', self.name_input)
        self.board_combo = QComboBox()
        for board in boards:
            self.board_combo.addItem(board['name'])
        layout.addRow('Board:', self.board_combo)
        self.task_combo = QComboBox()
        for task_name in tasks:
            self.task_combo.addItem(task_name)
        layout.addRow('Task:', self.task_combo)
        self.subject_list = QListWidget()
        for subject in subjects:
            item = QListWidgetItem(subject['name'])
            item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
            item.setCheckState(Qt.Unchecked)
            self.subject_list.addItem(item)
        layout.addRow('Subjects:', self.subject_list)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addRow(buttons)
        self.setLayout(layout)

    def setup_name(self):
        return self.name_input.text().strip()

    def board_name(self):
        return self.board_combo.currentText()

    def task_name(self):
        return self.task_combo.currentText()

    def selected_subjects(self):
        return [self.subject_list.item(i).text() for i in range(self.subject_list.count())
                if self.subject_list.item(i).checkState() == Qt.Checked]


# --- main window ---------------------------------------------------------------------------------

class SessionRunnerWindow(QWidget):
    def __init__(self):
        super(SessionRunnerWindow, self).__init__()
        self.setWindowTitle('Run Session -- AuditoryEvidenceAccum')

        self.launcher = TaskLauncher()
        self._stage_configs = discover_stage_configs()
        self._users = discover_users()
        self._watched_csv_path = None
        self._feeder = None
        self._stop_deadline = None

        self._build_ui()
        self._populate_stage_combo()
        self._run_task_python_check()

        self._session_timer = QTimer(self)
        self._session_timer.timeout.connect(self._on_session_poll)
        self._session_timer.start(VAR_SESSION_POLL_MS)

        self._plots_timer = QTimer(self)
        self._plots_timer.timeout.connect(self._on_plots_poll)
        self._plots_timer.start(VAR_PLOTS_POLL_MS)

        self._camera_timer = QTimer(self)
        self._camera_timer.timeout.connect(self._on_camera_poll)
        self._camera_timer.start(VAR_CAMERA_POLL_MS)

        self._stop_timer = QTimer(self)
        self._stop_timer.timeout.connect(self._on_stop_check)

    # --- UI ------------------------------------------------------------------------------------

    def _build_ui(self):
        outer = QVBoxLayout()

        # 1/3 of the actual screen's own width, queried live (same "don't hardcode a guess, ask
        # the real screen" convention wheel_shaping_plots.py's own _capped_figsize() already uses)
        # rather than a fixed pixel guess that would look right on one monitor and cramped/
        # oversized on another. Drives both the sidebar's own width and the camera preview's, so
        # the preview actually uses the space rather than sitting small in a now-wide sidebar.
        screen_width_px = QApplication.desktop().screenGeometry().width()
        sidebar_width_px = screen_width_px // 3
        self._camera_preview_width_px = (sidebar_width_px - 60) // 2

        self.task_python_warning_label = QLabel('')
        self.task_python_warning_label.setWordWrap(True)
        self.task_python_warning_label.setStyleSheet(
            'background-color: #ffcccc; color: #660000; padding: 6px; font-weight: bold;')
        self.task_python_warning_label.hide()
        outer.addWidget(self.task_python_warning_label)

        # Slim top strip -- reference info glanced at occasionally, not acted on repeatedly, so it
        # doesn't need a full group box competing for space with the sidebar below.
        top_widget = QWidget()
        top_layout = QVBoxLayout()
        top_layout.setContentsMargins(6, 4, 6, 4)
        top_layout.setSpacing(2)
        data_row = QHBoxLayout()
        data_row.addWidget(QLabel('<b>Session data:</b>'))
        self.session_data_label = QLabel('(none watched -- Start Task or Watch Latest Session)')
        self.session_data_label.setWordWrap(True)
        data_row.addWidget(self.session_data_label, stretch=1)
        top_layout.addLayout(data_row)
        log_row = QHBoxLayout()
        log_row.addWidget(QLabel('<b>Training log:</b>'))
        self.training_log_label = QLabel(_OUTPUT_XLSX)
        self.training_log_label.setWordWrap(True)
        log_row.addWidget(self.training_log_label, stretch=1)
        self.update_log_button = QPushButton('Update Training Log')
        self.update_log_button.clicked.connect(self._on_update_training_log)
        log_row.addWidget(self.update_log_button)
        top_layout.addLayout(log_row)
        self.update_log_status_label = QLabel('')
        self.update_log_status_label.setWordWrap(True)
        top_layout.addWidget(self.update_log_status_label)
        top_widget.setLayout(top_layout)
        outer.addWidget(top_widget)

        main_row = QHBoxLayout()

        # --- left sidebar: fixed, comfortable width -- a form-heavy control column doesn't need
        # to grow with the window, and looks messier if it does. Scrollable so nothing clips
        # regardless of screen height. Ordered top-to-bottom to match how these are actually used:
        # launch a task first, monitor it while it runs, fill in weight/notes last. -------------
        left_content = QWidget()
        left_layout = QVBoxLayout()

        launch_group = QGroupBox('Launch task')
        launch_form = QFormLayout()
        stage_row = QHBoxLayout()
        self.stage_combo = QComboBox()
        self.stage_combo.currentIndexChanged.connect(self._on_stage_changed)
        stage_row.addWidget(self.stage_combo, stretch=1)
        self.new_setup_button = QPushButton('New Setup...')
        self.new_setup_button.clicked.connect(self._on_new_setup)
        stage_row.addWidget(self.new_setup_button)
        launch_form.addRow('Stage:', stage_row)
        subject_row = QHBoxLayout()
        self.subject_combo = QComboBox()
        subject_row.addWidget(self.subject_combo, stretch=1)
        self.new_subject_button = QPushButton('New Subject...')
        self.new_subject_button.clicked.connect(self._on_new_subject)
        subject_row.addWidget(self.new_subject_button)
        launch_form.addRow('Subject:', subject_row)
        self.user_combo = QComboBox()
        for user in self._users:
            self.user_combo.addItem(user['name'], user)
        launch_form.addRow('User:', self.user_combo)
        buttons_row = QHBoxLayout()
        self.start_button = QPushButton('Start Task')
        self.start_button.clicked.connect(self._on_start_task)
        buttons_row.addWidget(self.start_button)
        self.stop_button = QPushButton('Stop Task')
        self.stop_button.clicked.connect(self._on_stop_task)
        self.stop_button.setEnabled(False)
        buttons_row.addWidget(self.stop_button)
        launch_form.addRow(buttons_row)
        self.watch_latest_button = QPushButton('Watch Latest Session')
        self.watch_latest_button.setToolTip(
            'Attach to whichever session is currently newest on disk -- e.g. one started from the '
            'real PyBpod GUI instead of this dashboard. The window opens blank; nothing is '
            'auto-watched.')
        self.watch_latest_button.clicked.connect(self._on_watch_latest)
        launch_form.addRow(self.watch_latest_button)
        self.launch_status_label = QLabel('Not running.')
        self.launch_status_label.setWordWrap(True)
        launch_form.addRow(self.launch_status_label)
        launch_group.setLayout(launch_form)
        left_layout.addWidget(launch_group)

        self.summary_label = QLabel('No session watched yet.')
        self.summary_label.setWordWrap(True)
        left_layout.addWidget(self.summary_label)

        camera_group = QGroupBox('Camera')
        camera_layout = QVBoxLayout()
        self.camera_label = QLabel('No session watched yet.')
        self.camera_label.setAlignment(Qt.AlignCenter)
        self.camera_label.setMinimumHeight(self._camera_preview_width_px * 3 // 4)
        self.camera_label.setStyleSheet('background-color: #222; color: #aaa;')
        camera_layout.addWidget(self.camera_label)
        camera_group.setLayout(camera_layout)
        left_layout.addWidget(camera_group)

        wheel_group = QGroupBox('Wheel position (session time; dotted = threshold, '
                                 'light vertical = trial start)')
        wheel_layout = QVBoxLayout()
        from matplotlib.figure import Figure
        from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg
        self._wheel_fig = Figure(figsize=(6, 3.6), constrained_layout=True)
        self._wheel_ax = self._wheel_fig.add_subplot(111)
        self._wheel_ax.set_xlabel('session time (s)', fontsize=8)
        self._wheel_ax.set_ylabel('wheel pos (deg)', fontsize=8)
        self._wheel_ax.tick_params(labelsize=7)
        self._wheel_canvas = FigureCanvasQTAgg(self._wheel_fig)
        self._wheel_canvas.setMinimumHeight(260)
        self._wheel_canvas.setMaximumHeight(340)
        wheel_layout.addWidget(self._wheel_canvas)
        wheel_group.setLayout(wheel_layout)
        left_layout.addWidget(wheel_group)

        weight_group = QGroupBox('Weight / notes (writes to training_log.xlsx)')
        weight_form = QFormLayout()
        self.dob_input = QLineEdit()
        weight_form.addRow('DOB:', self.dob_input)
        self.sex_input = QLineEdit()
        weight_form.addRow('Sex:', self.sex_input)
        self.strain_input = QLineEdit()
        weight_form.addRow('Strain:', self.strain_input)
        self.baseline_weight_input = QLineEdit()
        weight_form.addRow('Baseline weight (g):', self.baseline_weight_input)
        self.weight_before_input = QDoubleSpinBox()
        self.weight_before_input.setRange(0, 100)
        self.weight_before_input.setDecimals(1)
        weight_form.addRow('Weight before (g):', self.weight_before_input)
        self.weight_after_input = QDoubleSpinBox()
        self.weight_after_input.setRange(0, 100)
        self.weight_after_input.setDecimals(1)
        weight_form.addRow('Weight after (g):', self.weight_after_input)
        self.hand_watered_check = QCheckBox('Water given by hand after task')
        weight_form.addRow(self.hand_watered_check)
        self.notes_input = QTextEdit()
        self.notes_input.setMaximumHeight(80)
        weight_form.addRow('Notes:', self.notes_input)
        self.save_weight_button = QPushButton('Save to Training Log')
        self.save_weight_button.clicked.connect(self._on_save_weight)
        weight_form.addRow(self.save_weight_button)
        self.save_status_label = QLabel('')
        weight_form.addRow(self.save_status_label)
        weight_group.setLayout(weight_form)
        left_layout.addWidget(weight_group)

        left_layout.addStretch(1)
        left_content.setLayout(left_layout)

        left_scroll = QScrollArea()
        left_scroll.setWidgetResizable(True)
        left_scroll.setWidget(left_content)
        left_scroll.setFixedWidth(sidebar_width_px)
        left_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        main_row.addWidget(left_scroll, stretch=0)

        # --- main area (right, fills all remaining space): the full WheelShapingPlots outcome-
        # panel set, embedded directly (see PlotFeeder's own embed=True) instead of letting it
        # open as a separate popup window. Taller than most screens (esp. Stage 3/4's 6-row
        # mosaic), so it sits inside a scroll area. A fixed-width sidebar + an expanding main area
        # gives this the dominant share of the window on any real screen, without the awkward
        # competition a 50/50 split creates against a form column that doesn't need that much
        # width. -----------------------------------------------------------------------------
        self.plots_scroll = QScrollArea()
        self.plots_scroll.setWidgetResizable(True)
        self.plots_container = QWidget()
        self.plots_container_layout = QVBoxLayout()
        self.plots_placeholder_label = QLabel(
            'No session watched yet.\n\nClick "Start Task" to launch and watch a new session, or '
            '"Watch Latest Session" to attach to one already running (e.g. started from the real '
            'PyBpod GUI).\n\nOnce trials start arriving, the full outcome-panel set (raster, '
            'accuracy, sidebias, psychometric, lick panels, ...) will appear here.')
        self.plots_placeholder_label.setAlignment(Qt.AlignCenter)
        self.plots_placeholder_label.setWordWrap(True)
        self.plots_placeholder_label.setStyleSheet('color: #666; font-style: italic; padding: 40px;')
        self.plots_container_layout.addWidget(self.plots_placeholder_label)
        self.plots_container.setLayout(self.plots_container_layout)
        self.plots_scroll.setWidget(self.plots_container)
        main_row.addWidget(self.plots_scroll, stretch=1)

        outer.addLayout(main_row, stretch=1)
        self.setLayout(outer)

    def _populate_stage_combo(self):
        self.stage_combo.clear()
        for cfg in self._stage_configs:
            label = '{0} ({1})'.format(cfg['experiment'], cfg['task_name'])
            self.stage_combo.addItem(label, cfg)
        self._on_stage_changed()

    def _on_stage_changed(self):
        self.subject_combo.clear()
        cfg = self.stage_combo.currentData()
        if cfg is None:
            return
        for subject in cfg['subjects']:
            self.subject_combo.addItem(subject)

    # --- task-python self-check / training log / new subject+setup -----------------------------

    def _run_task_python_check(self):
        if task_python_can_import_pybpodapi():
            self.task_python_warning_label.hide()
            return
        self.task_python_warning_label.setText(
            'WARNING: task launches will fail -- "pybpodapi" is not importable via "{0}". Edit '
            'VAR_TASK_PYTHON_EXE at the top of run_session.py to point at your pybpod-environment\'s'
            ' own python.exe.'.format(VAR_TASK_PYTHON_EXE))
        self.task_python_warning_label.show()

    def _on_update_training_log(self):
        buf = io.StringIO()
        try:
            with contextlib.redirect_stdout(buf):
                build_training_log.build_workbook()
        except Exception as err:
            print(traceback.format_exc(), flush=True)
            QMessageBox.critical(self, 'Update failed', str(err))
            return
        output = buf.getvalue().strip()
        if output:
            print(output, flush=True)
        last_line = output.split('\n')[-1] if output else 'Done.'
        self.update_log_status_label.setText(last_line)

    def _on_new_subject(self):
        dialog = NewSubjectDialog(self)
        if dialog.exec_() != QDialog.Accepted:
            return
        name = dialog.subject_name()
        try:
            create_subject(name)
        except Exception as err:
            QMessageBox.critical(self, 'Could not create subject', str(err))
            return
        QMessageBox.information(
            self, 'Subject created',
            'Created subject "{0}". Attach it to a Setup via "New Setup..." (or an existing one) '
            'to use it for a task.'.format(name))

    def _on_new_setup(self):
        boards = discover_boards()
        if not boards:
            QMessageBox.warning(self, 'No boards',
                                 'No boards found under boards/ -- create one via the PyBpod GUI '
                                 'first.')
            return
        dialog = NewSetupDialog(boards, discover_all_tasks(), discover_subjects(), self)
        if dialog.exec_() != QDialog.Accepted:
            return
        name = dialog.setup_name()
        try:
            create_setup(name, dialog.board_name(), dialog.task_name(), dialog.selected_subjects())
        except Exception as err:
            QMessageBox.critical(self, 'Could not create setup', str(err))
            return
        self._stage_configs = discover_stage_configs()
        self._populate_stage_combo()
        QMessageBox.information(self, 'Setup created',
                                 'Created experiment/setup "{0}".'.format(name))

    # --- launch / stop ---------------------------------------------------------------------------

    def _on_start_task(self):
        cfg = self.stage_combo.currentData()
        subject = self.subject_combo.currentText()
        user = self.user_combo.currentData()
        if cfg is None or not subject or user is None:
            QMessageBox.warning(self, 'Missing selection',
                                 'Choose a stage, subject, and user first.')
            return
        try:
            session_path = self.launcher.start(cfg, subject, user)
        except Exception as err:
            print(traceback.format_exc(), flush=True)
            QMessageBox.critical(self, 'Launch failed', str(err))
            return
        self.launch_status_label.setText('Running: {0}'.format(session_path))
        self.start_button.setEnabled(False)
        self.stop_button.setEnabled(True)
        self._watch_session(os.path.join(
            session_path, os.path.basename(session_path) + '.csv'))
        QTimer.singleShot(VAR_QUICK_DEATH_CHECK_MS, self._on_check_quick_death)

    def _on_check_quick_death(self):
        """ If the task process already died within VAR_QUICK_DEATH_CHECK_MS of launching, surface
        its captured stderr directly in the UI instead of leaving it silently sitting in a log file
        inside the new session folder -- see VAR_TASK_PYTHON_EXE's own docstring for the failure
        mode this exists to catch. """
        proc = self.launcher.proc
        if proc is None or proc.poll() is None:
            return   # still running (or a newer launch has already replaced this one)
        session_path = self.launcher.session_path
        tail = ''
        if session_path:
            stderr_path = os.path.join(session_path, 'dashboard_launch_stderr.log')
            try:
                with open(stderr_path, 'r', encoding='utf-8', errors='replace') as f:
                    tail = f.read().strip()[-4000:]
            except Exception:
                pass
        QMessageBox.critical(
            self, 'Task exited immediately',
            'The task process exited right after launching (exit code {0}).\n\n{1}'.format(
                proc.returncode, tail or '(no stderr captured -- check {0})'.format(session_path)))
        self._on_task_finished()

    def _on_stop_task(self):
        if not self.launcher.is_running:
            return
        self.launcher.stop(graceful=True)
        self.stop_button.setEnabled(False)
        self.launch_status_label.setText('Stopping (waiting up to {0:.0f}s for a clean '
                                          'finish)...'.format(VAR_STOP_GRACE_S))
        self._stop_deadline = time.time() + VAR_STOP_GRACE_S
        self._stop_timer.start(500)

    def _on_stop_check(self):
        if not self.launcher.is_running:
            self._stop_timer.stop()
            self._on_task_finished()
            return
        if time.time() >= self._stop_deadline:
            self.launcher.force_kill()

    def _on_task_finished(self):
        self.start_button.setEnabled(True)
        self.stop_button.setEnabled(False)
        self.launch_status_label.setText('Finished.')

    # --- session watching --------------------------------------------------------------------

    def _watch_session(self, csv_path):
        self._watched_csv_path = csv_path
        self._feeder = None
        self.session_data_label.setText(csv_path)
        self.summary_label.setText('Watching -- waiting for trial data...')

        # Reset the camera/wheel/outcome panels immediately -- otherwise the PREVIOUS session's
        # own stale frame/plot would stay visible until the newly watched session's own first
        # frame/trial arrives.
        self.camera_label.setPixmap(QPixmap())
        self.camera_label.setText('No camera frame yet.')
        self._wheel_ax.cla()
        self._wheel_ax.set_xlabel('session time (s)', fontsize=8)
        self._wheel_ax.set_ylabel('wheel pos (deg)', fontsize=8)
        self._wheel_ax.tick_params(labelsize=7)
        self._wheel_canvas.draw_idle()
        self.plots_placeholder_label.setText(
            'Watching this session -- the full outcome-panel set will appear here once it has at '
            'least one trial.')
        self._embed_plots_canvas(self.plots_placeholder_label)

    def _on_session_poll(self):
        # Deliberately does NOT auto-watch whatever session is newest on disk -- the window opens
        # to a blank slate; only Start Task or "Watch Latest Session" ever picks a session to
        # watch. This only tracks whether a task THIS dashboard itself launched has finished.
        if self.launcher.is_running and self.launcher.proc.poll() is not None:
            self._on_task_finished()

    def _on_watch_latest(self):
        newest = find_newest_session_csv()
        if newest is None:
            QMessageBox.information(self, 'No sessions found',
                                     'No session CSVs found on disk yet.')
            return
        self._watch_session(newest)

    # --- live plots ----------------------------------------------------------------------------

    def _on_plots_poll(self):
        if self._watched_csv_path is None or not os.path.exists(self._watched_csv_path):
            return
        try:
            info, session_vals, trials = session_csv_parser.parse_session_csv(self._watched_csv_path)
        except Exception:
            print(traceback.format_exc(), flush=True)
            return
        protocol_name = info.get('PROTOCOL-NAME')
        real = session_csv_parser.real_trials(trials)

        summary = None
        if protocol_name in session_csv_parser.PROTOCOL_CONFIG:
            try:
                summary = build_training_log.summarize_session(self._watched_csv_path)
            except Exception:
                print(traceback.format_exc(), flush=True)
        if summary is not None:
            self.summary_label.setText(
                'Subject: {subject}  |  Protocol: {protocol}  |  Trials: {trial_count}  |  '
                'Rewarded: {reward_count}  |  Licks: {lick_count}'.format(**summary))
            self._current_subject = summary['subject']
            self._current_session_started = summary['session_started']

        self._redraw_wheel_plot(real, protocol_name, session_vals)

        if protocol_name in _DASHBOARD_PROTOCOL_INFO and real:
            if self._feeder is None or self._feeder.protocol_name != protocol_name:
                self._feeder = PlotFeeder(protocol_name, real[0]['vals'], session_vals)
                self._embed_plots_canvas(self._feeder.plots.canvas)
            try:
                self._feeder.feed_new_trials(real, session_ended=('SESSION-ENDED' in info))
            except Exception:
                print(traceback.format_exc(), flush=True)

    def _redraw_wheel_plot(self, real, protocol_name, session_vals):
        """ Small cla()+rebuild every poll -- cheap at this cadence/session length (same "small
        panel, cheap redraw" convention already used elsewhere for low-cardinality panels; only
        the FULL WheelShapingPlots raster needs persistent artists, since that one redraws every
        single trial rather than every ~2s poll). """
        xs, ys = parse_wheel_position_series(real)
        starts, segments = build_trial_markers(real, protocol_name, session_vals)
        self._wheel_ax.cla()
        for x in starts:
            self._wheel_ax.axvline(x, color='#dddddd', linewidth=0.6, zorder=0)
        for x_start, x_end, threshold in segments:
            self._wheel_ax.plot([x_start, x_end], [threshold, threshold], color='black',
                                 linestyle=':', linewidth=0.9, zorder=1)
            self._wheel_ax.plot([x_start, x_end], [-threshold, -threshold], color='black',
                                 linestyle=':', linewidth=0.9, zorder=1)
        self._wheel_ax.plot(xs, ys, linewidth=0.8, color='tab:blue', zorder=2)
        self._wheel_ax.axhline(0, color='#888888', linewidth=0.5, zorder=0)
        self._wheel_ax.set_xlabel('session time (s)', fontsize=8)
        self._wheel_ax.set_ylabel('wheel pos (deg)', fontsize=8)
        self._wheel_ax.tick_params(labelsize=7)
        if xs or segments:
            all_x = xs + [s[0] for s in segments] + [s[1] for s in segments]
            all_y = ys + [s[2] for s in segments] + [-s[2] for s in segments]
            self._wheel_ax.set_xlim(min(all_x), max(all_x) + 1)
            y_span = max([abs(y) for y in all_y] + [5.0]) * 1.1
            self._wheel_ax.set_ylim(-y_span, y_span)
        self._wheel_canvas.draw_idle()

    def _embed_plots_canvas(self, canvas):
        """ Swaps the currently-shown WheelShapingPlots canvas (or the placeholder label, the
        first time) for a new one -- happens whenever a new PlotFeeder is created, i.e. a
        different session or a different protocol/stage is now being watched. """
        old = self.plots_container_layout.takeAt(0)
        if old is not None:
            widget = old.widget()
            if widget is not None:
                widget.setParent(None)
        self.plots_container_layout.addWidget(canvas)

    # --- camera --------------------------------------------------------------------------------

    def _on_camera_poll(self):
        if self._watched_csv_path is None:
            return
        session_dir = os.path.dirname(self._watched_csv_path)
        try:
            frame = grab_latest_frame(session_dir)
        except Exception:
            print(traceback.format_exc(), flush=True)
            return
        if frame is None:
            return
        rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        rgb = np.ascontiguousarray(rgb)
        h, w, _ch = rgb.shape
        qimage = QImage(rgb.data, w, h, 3 * w, QImage.Format_RGB888)
        pixmap = QPixmap.fromImage(qimage).scaledToWidth(
            self._camera_preview_width_px, Qt.SmoothTransformation)
        self.camera_label.setPixmap(pixmap)

    # --- weight / notes save -------------------------------------------------------------------

    def _on_save_weight(self):
        subject = self.subject_combo.currentText() or getattr(self, '_current_subject', None)
        session_started = getattr(self, '_current_session_started', None)
        if not subject:
            QMessageBox.warning(self, 'No subject', 'Choose a subject (or watch a session) first.')
            return
        if not session_started:
            QMessageBox.warning(self, 'No session', 'No watched session to attach this entry to '
                                                      'yet -- start or wait for a session first.')
            return
        header_fields = [self.dob_input.text().strip(), self.sex_input.text().strip(),
                          self.strain_input.text().strip(), self.baseline_weight_input.text().strip()]
        try:
            found = save_manual_fields(
                subject, session_started,
                weight_g=(self.weight_before_input.value() or None),
                weight_after_g=(self.weight_after_input.value() or None),
                hand_watered=self.hand_watered_check.isChecked(),
                notes=self.notes_input.toPlainText().strip(),
                header_fields=header_fields)
        except Exception as err:
            print(traceback.format_exc(), flush=True)
            QMessageBox.critical(self, 'Save failed', str(err))
            return
        if found:
            self.save_status_label.setText('Saved.')
        else:
            self.save_status_label.setText('Header fields saved -- session row not found yet '
                                            '(try again once the session has at least one trial).')

    # --- shutdown --------------------------------------------------------------------------------

    def closeEvent(self, event):
        if self.launcher.is_running:
            reply = QMessageBox.question(
                self, 'Task still running',
                'A task launched from this dashboard is still running. Stop it before closing?',
                QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel)
            if reply == QMessageBox.Cancel:
                event.ignore()
                return
            if reply == QMessageBox.Yes:
                self.launcher.stop(graceful=True)
        event.accept()


if __name__ == '__main__':
    app = QApplication(sys.argv)
    window = SessionRunnerWindow()
    window.showMaximized()
    sys.exit(app.exec_())
