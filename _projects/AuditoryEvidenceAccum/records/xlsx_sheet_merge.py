"""Merge two versions of training_log.xlsx sheet by sheet (one sheet per mouse).

git can't merge a binary .xlsx, but this workbook doesn't need a cell-level merge: each sheet
belongs to exactly one mouse, and each mouse is trained on exactly one computer (see
build_training_log.owned_subjects()). So when two computers have both committed the workbook,
the merged result is simply: every sheet from THEIR version, except the sheets this computer owns,
which come from OUR version. Sheets only one side has are kept as-is.

Used by sync_training_log.py; the merged bytes are what gets committed.
"""
import copy
import io

import openpyxl


def _copy_sheet(src, dst):
    """openpyxl can't move a worksheet between workbooks, so copy what build_training_log.py
    actually uses: values/formulas, cell styles, merged ranges, column widths/hidden flags, row
    heights, freeze panes, tab color."""
    for row in src.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            new = dst.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new.font = copy.copy(cell.font)
                new.fill = copy.copy(cell.fill)
                new.border = copy.copy(cell.border)
                new.alignment = copy.copy(cell.alignment)
                new.protection = copy.copy(cell.protection)
                new.number_format = cell.number_format
    for merged in src.merged_cells.ranges:
        dst.merge_cells(str(merged))
    for key, dim in src.column_dimensions.items():
        dst.column_dimensions[key].width = dim.width
        dst.column_dimensions[key].hidden = dim.hidden
    for key, dim in src.row_dimensions.items():
        if dim.height is not None:
            dst.row_dimensions[key].height = dim.height
    dst.freeze_panes = src.freeze_panes
    dst.sheet_properties.tabColor = src.sheet_properties.tabColor


def merge_workbooks(theirs_bytes, ours_bytes, our_sheet_names):
    """Returns the merged workbook as bytes: `theirs`, with every sheet named in
    `our_sheet_names` (plus any sheet only `ours` has) replaced by/added from `ours`."""
    theirs = openpyxl.load_workbook(io.BytesIO(theirs_bytes))
    ours = openpyxl.load_workbook(io.BytesIO(ours_bytes))
    for name in ours.sheetnames:
        if name in theirs.sheetnames:
            if name not in our_sheet_names:
                continue
            index = theirs.sheetnames.index(name)
            theirs.remove(theirs[name])
        else:
            index = len(theirs.sheetnames)
        _copy_sheet(ours[name], theirs.create_sheet(name, index))
    out = io.BytesIO()
    theirs.save(out)
    return out.getvalue()
